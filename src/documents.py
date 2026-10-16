"""Document pipeline for retrieval-augmented generation (RAG).

    file path -> read text -> split into chunks -> embed each chunk
              -> store vectors in Chroma (data/chroma)
              -> store one row per document in the `documents` table (agent.db)

    question  -> embed the question -> nearest chunks -> the agent reads them

Heavy libraries (chromadb, pypdf, python-docx) are imported only when needed,
so importing this module is cheap and `list_all()` / `names_for_prompt()` never
need Ollama or Chroma.

The module-level functions at the bottom (add_file, list_all, remove, search)
use one default store. They are also what the /doc chat command calls.
"""
import hashlib
import uuid
from datetime import datetime
from pathlib import Path

from src import config
from src.memory import conn as default_conn

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".csv", ".json", ".log", ".py"}
SUPPORTED = sorted(TEXT_EXTENSIONS | {".pdf", ".docx"})
EMBED_BATCH = 16


# =====================================================================
# Reading and chunking (pure functions, easy to test)
# =====================================================================
def read_pages(path: Path) -> list[tuple[int | None, str]]:
    """Return [(page_number_or_None, text), ...]. Only PDFs have page numbers."""
    ext = path.suffix.lower()
    if ext in TEXT_EXTENSIONS:
        return [(None, path.read_text(encoding="utf-8", errors="replace"))]
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("PDF support needs pypdf: pip install pypdf")
        reader = PdfReader(str(path))
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError("Word support needs python-docx: pip install python-docx")
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:                           # tables hold text too
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return [(None, "\n".join(parts))]
    raise ValueError(f"Unsupported file type '{ext}'. Supported: {', '.join(SUPPORTED)}")


def split_pages(pages, chunk_size: int, overlap: int) -> list[tuple[int | None, str]]:
    """Split each page into overlapping chunks, keeping the page number."""
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=overlap)
    chunks = []
    for page, text in pages:
        for piece in splitter.split_text(text):
            if piece.strip():
                chunks.append((page, piece.strip()))
    return chunks


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


# =====================================================================
# The store
# =====================================================================
class DocumentStore:
    def __init__(self, conn=None, chroma_dir=None, embeddings=None):
        """conn: sqlite connection for the `documents` table.
        chroma_dir: folder of the vector database.
        embeddings: object with embed_documents(list) and embed_query(str).
        Created lazily from config if omitted. Pass your own to use a separate store."""
        self.conn = conn or default_conn
        self.chroma_dir = Path(chroma_dir or config.CHROMA_DIR)
        self._embeddings = embeddings
        self._collection = None
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS documents (
                id       TEXT PRIMARY KEY,
                name     TEXT NOT NULL,
                path     TEXT NOT NULL,
                chunks   INTEGER NOT NULL,
                sha256   TEXT NOT NULL UNIQUE,
                added_at TEXT NOT NULL
            )
        """)
        self.conn.commit()

    # ---------- lazy resources ----------
    def _col(self):
        if self._collection is None:
            import chromadb
            from chromadb.config import Settings
            self.chroma_dir.mkdir(parents=True, exist_ok=True)
            client = chromadb.PersistentClient(
                path=str(self.chroma_dir), settings=Settings(anonymized_telemetry=False))
            try:                                           # newer Chroma
                self._collection = client.get_or_create_collection(
                    "documents", configuration={"hnsw": {"space": "cosine"}})
            except (TypeError, ValueError):                # older Chroma
                self._collection = client.get_or_create_collection(
                    "documents", metadata={"hnsw:space": "cosine"})
        return self._collection

    def _embed(self):
        if self._embeddings is None:
            self._embeddings = config.get_embeddings()
        return self._embeddings

    def _embed_texts(self, texts: list[str], prefix: str, on_progress=None) -> list[list[float]]:
        vectors = []
        try:
            for i in range(0, len(texts), EMBED_BATCH):
                batch = [prefix + t for t in texts[i:i + EMBED_BATCH]]
                vectors.extend(self._embed().embed_documents(batch))
                if on_progress:
                    on_progress(min(i + EMBED_BATCH, len(texts)), len(texts))
        except Exception as e:
            raise RuntimeError(
                f"Embedding failed ({type(e).__name__}: {e}). Is Ollama running and the model "
                f"pulled? Run: ollama pull {config.EMBED_MODEL}") from e
        return vectors

    # ---------- listing and lookup ----------
    def list_all(self) -> list[dict]:
        rows = self.conn.execute(
            "SELECT id, name, path, chunks, added_at FROM documents ORDER BY added_at DESC").fetchall()
        return [dict(zip(("id", "name", "path", "chunks", "added_at"), r)) for r in rows]

    def count(self) -> int:
        return self.conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0]

    def find(self, ref: str) -> dict | None:
        """Find by exact id, id prefix or name (case-insensitive).
        Raises ValueError if the reference matches more than one document."""
        ref = ref.strip().strip('"').strip("'").lower()
        docs = self.list_all()
        for d in docs:
            if d["id"] == ref:
                return d
        matches = [d for d in docs if d["id"].startswith(ref) or d["name"].lower() == ref]
        if len(matches) > 1:
            raise ValueError(f"'{ref}' matches several documents; use the id instead.")
        return matches[0] if matches else None

    def names_for_prompt(self, limit: int = 15) -> str:
        names = [d["name"] for d in self.list_all()[:limit]]
        return ", ".join(names)

    # ---------- adding ----------
    def add_file(self, path, on_progress=None) -> dict:
        p = Path(str(path).strip().strip('"').strip("'")).expanduser()
        if not p.is_file():
            raise FileNotFoundError(f"No such file: {p}")

        digest = _sha256(p)
        row = self.conn.execute("SELECT id, name FROM documents WHERE sha256 = ?", (digest,)).fetchone()
        if row:
            raise ValueError(f"This file is already indexed as '{row[1]}' (id {row[0]}).")

        chunks = split_pages(read_pages(p), config.CHUNK_SIZE, config.CHUNK_OVERLAP)
        if not chunks:
            raise ValueError("No text could be extracted. Scanned PDFs (images of text) "
                             "are not supported; they need OCR first.")
        if len(chunks) > config.MAX_CHUNKS_PER_FILE:
            raise ValueError(f"File too large ({len(chunks)} chunks, limit "
                             f"{config.MAX_CHUNKS_PER_FILE}).")

        texts = [c[1] for c in chunks]
        vectors = self._embed_texts(texts, config.EMBED_DOC_PREFIX, on_progress)

        doc_id = uuid.uuid4().hex[:8]
        self._col().add(
            ids=[f"{doc_id}:{i}" for i in range(len(chunks))],
            documents=texts,
            embeddings=vectors,
            metadatas=[{"doc_id": doc_id, "name": p.name, "page": page or 0, "chunk": i}
                       for i, (page, _) in enumerate(chunks)],
        )
        # Record the document only after its vectors are safely stored
        self.conn.execute(
            "INSERT INTO documents (id, name, path, chunks, sha256, added_at) VALUES (?, ?, ?, ?, ?, ?)",
            (doc_id, p.name, str(p.resolve()), len(chunks), digest, datetime.now().isoformat(timespec="seconds")))
        self.conn.commit()
        return {"id": doc_id, "name": p.name, "chunks": len(chunks)}

    # ---------- removing ----------
    def remove(self, ref: str) -> bool:
        doc = self.find(ref)
        if doc is None:
            return False
        self._col().delete(where={"doc_id": doc["id"]})
        self.conn.execute("DELETE FROM documents WHERE id = ?", (doc["id"],))
        self.conn.commit()
        return True

    # ---------- searching ----------
    def search(self, query: str, k: int | None = None, max_distance: float | None = None) -> list[dict]:
        """Return the most similar chunks: [{text, name, page, score, distance}, ...].
        Chunks farther than max_distance are dropped as irrelevant."""
        if self.count() == 0 or not query.strip():
            return []
        k = k or config.TOP_K
        limit = config.MAX_DISTANCE if max_distance is None else max_distance
        col = self._col()
        n = col.count()
        if n == 0:
            return []
        vector = self._embed_texts([query], config.EMBED_QUERY_PREFIX)[0]
        res = col.query(query_embeddings=[vector], n_results=min(k, n),
                        include=["documents", "metadatas", "distances"])
        hits = []
        for text, meta, dist in zip(res["documents"][0], res["metadatas"][0], res["distances"][0]):
            if dist <= limit:
                hits.append({"text": text, "name": meta["name"], "page": meta.get("page") or None,
                             "score": round(1 - dist, 3), "distance": round(dist, 3)})
        return hits


# =====================================================================
# Default store and the interface the /doc command uses
# =====================================================================
store = DocumentStore()


def add_file(path, on_progress=None) -> dict:
    return store.add_file(path, on_progress)


def list_all() -> list[dict]:
    return store.list_all()


def remove(ref: str) -> bool:
    return store.remove(ref)


def search(query: str, k: int | None = None) -> list[dict]:
    return store.search(query, k)